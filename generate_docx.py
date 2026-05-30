from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_documentation():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Title Page
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run('MINOR PROJECT TOPIC')
    title_run.bold = True
    title_run.font.size = Pt(16)
    
    doc.add_paragraph()  # Empty line
    
    project_para = doc.add_paragraph()
    project_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project_run = project_para.add_run('InstaHub')
    project_run.bold = True
    project_run.font.size = Pt(14)
    
    doc.add_paragraph()
    
    session_para = doc.add_paragraph()
    session_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    session_run = session_para.add_run('SESSION: 2025-26')
    session_run.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    submitted_para = doc.add_paragraph()
    submitted_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    submitted_run = submitted_para.add_run('SUBMITTED BY:')
    submitted_run.bold = True
    
    doc.add_paragraph('ADESH GUPTA', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('ROLL NUMBER:', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('17824407023', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    submitted_to_para = doc.add_paragraph()
    submitted_to_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    submitted_to_run = submitted_to_para.add_run('SUBMITTED TO:')
    submitted_to_run.bold = True
    
    doc.add_paragraph('VIVEK SIR', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('SIGNATURE: ________________', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Page break
    doc.add_page_break()
    
    # Acknowledgement
    heading = doc.add_heading('ACKNOWLEDGEMENT', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(
        'I would like to express my sincere gratitude to all those who supported and guided me throughout the successful completion of this minor project. First and foremost, I am deeply thankful to my project guide for their valuable supervision, insightful suggestions, and continuous encouragement, which greatly contributed to the quality and completion of this work. I also extend my heartfelt thanks to the faculty members of the Department of Computer Applications for providing the academic foundation and resources that enabled me to undertake this project.'
    )
    
    doc.add_paragraph(
        'I am grateful to my institution for offering a conducive learning environment and the necessary facilities for carrying out this study. My sincere appreciation is also due to my classmates and friends for their cooperation, constructive feedback, and moral support during the various stages of the project.'
    )
    
    doc.add_paragraph(
        'Finally, I would like to acknowledge the constant encouragement and understanding of my family, whose support has been an important source of motivation throughout my academic journey. Without the guidance, assistance, and goodwill of all these individuals, this project would not have been possible.'
    )
    
    # Page break
    doc.add_page_break()
    
    # Table of Contents
    toc_heading = doc.add_heading('TABLE OF CONTENTS', level=1)
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'TOPIC'
    hdr_cells[1].text = 'PAGE NO.'
    hdr_cells[2].text = "TEACHER'S SIGNATURE"
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].bold = True
    
    # Add rows
    topics = [
        ('1. INTRODUCTION', '4', ''),
        ('    1.1 Tech Stack', '4', ''),
        ('    1.2 Installation Steps', '4', ''),
        ('    1.3 How to Run the Project', '4', ''),
        ('2. REQUIREMENTS', '5', ''),
        ('    2.1 Hardware Requirements', '5', ''),
        ('    2.2 Software Requirements', '5', ''),
        ('3. Entity Relationship (ER) Diagram', '6', ''),
        ('4. CODE & UI SNAPSHOTS', '7-XX', ''),
        ('    4.1 Main Application (main.py)', '7', ''),
        ('    4.2 Configuration Files', '8', ''),
        ('    4.3 Authentication Module', '9', ''),
        ('    4.4 Controllers', '10', ''),
        ('    4.5 Frontend', '11', ''),
        ('5. TESTING', 'XX', ''),
        ('6. FUTURE ENHANCEMENT', 'XX', ''),
        ('7. CONCLUSION', 'XX', ''),
    ]
    
    for topic, page, sig in topics:
        row_cells = table.add_row().cells
        row_cells[0].text = topic
        row_cells[1].text = page
        row_cells[2].text = sig
    
    # Page break
    doc.add_page_break()
    
    # Introduction
    intro_heading = doc.add_heading('INTRODUCTION', level=1)
    intro_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(
        'Social media platforms have become an integral part of our daily lives, connecting people across the globe through shared content and interactions. InstaHub is an Instagram-like social media platform that allows users to share images, connect with others, and engage through likes, comments, and follows. This project was developed to provide a practical understanding of full-stack web development, combining modern backend technologies with responsive frontend design.'
    )
    
    doc.add_paragraph(
        'The InstaHub project demonstrates the implementation of a complete social media platform MVP, including user authentication, post management, social interactions, and a personalized feed system. This project helped me apply concepts learned during my BCA course, including RESTful API design, database management, authentication mechanisms, and frontend development. It also provided real-world experience in building a practical software application from start to finish.'
    )
    
    # Tech Stack
    tech_heading = doc.add_heading('1.1 Tech Stack:', level=2)
    tech_list = [
        'Backend: Python, FastAPI',
        'Database: MongoDB',
        'Frontend: HTML, CSS, JavaScript',
        'Authentication: JWT (JSON Web Tokens)',
        'Storage: Local file storage for images',
        'Tools/Editor: VS Code or any preferred IDE',
        'Browser: Chrome / Edge / Firefox'
    ]
    for item in tech_list:
        para = doc.add_paragraph(item, style='List Bullet')
    
    # Installation Steps
    install_heading = doc.add_heading('1.2 Installation Steps', level=2)
    install_steps = [
        'Install Python 3.8 or higher',
        'Install MongoDB and ensure it\'s running locally',
        'Install VS Code or any code editor',
        'Navigate to project folder',
        'Create a virtual environment (recommended): python -m venv venv',
        'Activate virtual environment:\n   • Windows: venv\\Scripts\\activate\n   • Linux/Mac: source venv/bin/activate',
        'Run pip install -r requirements.txt',
        'Create a .env file in the project root with:\n   MONGODB_URL=mongodb://localhost:27017\n   DATABASE_NAME=instahub\n   SECRET_KEY=your-secret-key-here-change-in-production'
    ]
    for i, step in enumerate(install_steps, 1):
        para = doc.add_paragraph(f'{i}. {step}', style='List Number')
    
    # How to Run
    run_heading = doc.add_heading('1.3 How to Run the Project', level=2)
    run_steps = [
        'Ensure MongoDB is running on your system',
        'Activate the virtual environment (if using one)',
        'Windows: Double-click start.bat or run: python -m uvicorn main:app --reload --host 0.0.0.0 --port 8000',
        'Linux/Mac: Run: uvicorn main:app --reload',
        'Open your browser and navigate to http://localhost:8000',
        'API documentation is available at http://localhost:8000/docs'
    ]
    for i, step in enumerate(run_steps, 1):
        para = doc.add_paragraph(f'{i}. {step}', style='List Number')
    
    # Page break
    doc.add_page_break()
    
    # Requirements
    req_heading = doc.add_heading('REQUIREMENT', level=1)
    req_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    hw_heading = doc.add_heading('2.1 Hardware Requirements:', level=2)
    hw_list = [
        'Personal computer or laptop',
        'Processor: Minimum Dual Core (recommended Core i3 or above)',
        'RAM: Minimum 2 GB (recommended 4 GB)',
        'Storage: Minimum 5 GB free disk space',
        'Internet connectivity (for web-based version or online deployment)'
    ]
    for item in hw_list:
        doc.add_paragraph(item, style='List Bullet')
    
    sw_heading = doc.add_heading('2.2 Software Requirements:', level=2)
    sw_list = [
        'Operating System: Windows / Linux / macOS',
        'Programming Language: Python 3.8 or higher',
        'Backend Framework: FastAPI',
        'Database: MongoDB',
        'Front-end technologies: HTML, CSS, JavaScript',
        'Code Editor or IDE: VS Code',
        'Web Browser: Chrome / Edge / Firefox',
        'Python Package Manager: pip'
    ]
    for item in sw_list:
        doc.add_paragraph(item, style='List Bullet')
    
    # Page break
    doc.add_page_break()
    
    # ER Diagram
    er_heading = doc.add_heading('Entity Relationship Diagram', level=1)
    er_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    er_para = doc.add_paragraph()
    er_run = er_para.add_run('[Please insert ER Diagram here]')
    er_run.bold = True
    er_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('The ER Diagram should represent the relationships between:')
    er_entities = ['Users', 'Posts', 'Comments', 'Likes', 'Follows', 'Notifications']
    for entity in er_entities:
        doc.add_paragraph(f'• {entity}', style='List Bullet')
    
    # Page break
    doc.add_page_break()
    
    # Code & UI Snapshots
    code_heading = doc.add_heading('CODE & UI SNAPSHOTS', level=1)
    code_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    main_heading = doc.add_heading('4.1 Main Application (main.py)', level=2)
    doc.add_paragraph('The main.py file serves as the entry point of the FastAPI application, defining all routes and middleware configurations.')
    
    # Code block for main.py
    code_text = """from fastapi import FastAPI, Depends, UploadFile, File, Form, HTTPException
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from typing import Optional
import os

from config import IMAGES_DIR, PROFILES_DIR
from database import db
from auth import get_current_user
from controllers import (
    auth_controller,
    user_controller,
    post_controller,
    feed_controller,
    notification_controller
)

# Create FastAPI app
app = FastAPI(title="InstaHub MVP", version="1.0.0")

# CORS middleware configuration
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"],
)

# Create upload directories
os.makedirs(IMAGES_DIR, exist_ok=True)
os.makedirs(PROFILES_DIR, exist_ok=True)

# Mount static files
app.mount("/uploads", StaticFiles(directory="uploads"), name="uploads")
app.mount("/frontend", StaticFiles(directory="frontend"), name="frontend")

# Root route
@app.get("/")
async def root():
    return FileResponse("frontend/index.html")

# Authentication routes
@app.post("/auth/register")
async def register(
    username: str = Form(...),
    email: str = Form(...),
    password: str = Form(...),
    bio: Optional[str] = Form(None),
    profile_image: Optional[UploadFile] = File(None)
):
    return await auth_controller.register_user(username, email, password, bio, profile_image)

@app.post("/auth/login")
async def login(username: str = Form(...), password: str = Form(...)):
    return await auth_controller.login_user(username, password)

# User routes
@app.get("/user/{user_id}")
async def get_user(user_id: str, current_user: dict = Depends(get_current_user)):
    return await user_controller.get_user_profile(user_id, current_user)

@app.post("/follow/{user_id}")
async def follow_user(user_id: str, current_user: dict = Depends(get_current_user)):
    return await user_controller.follow_unfollow_user(user_id, current_user)

# Post routes
@app.post("/posts")
async def create_post(
    image: UploadFile = File(...),
    caption: str = Form(""),
    current_user: dict = Depends(get_current_user)
):
    return await post_controller.create_post(image, caption, current_user)

@app.post("/posts/{post_id}/like")
async def like_post(post_id: str, current_user: dict = Depends(get_current_user)):
    return await post_controller.like_unlike_post(post_id, current_user)

@app.post("/posts/{post_id}/comment")
async def add_comment(
    post_id: str,
    comment_text: str = Form(...),
    current_user: dict = Depends(get_current_user)
):
    return await post_controller.add_comment(post_id, comment_text, current_user)

# Feed routes
@app.get("/feed")
async def get_feed(current_user: dict = Depends(get_current_user)):
    return await feed_controller.get_feed(current_user)

# Notification routes
@app.get("/notifications")
async def get_notifications(current_user: dict = Depends(get_current_user)):
    return await notification_controller.get_user_notifications(current_user)"""
    
    code_para = doc.add_paragraph(code_text)
    code_para.style = 'No Spacing'
    for run in code_para.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    
    # Config section
    config_heading = doc.add_heading('4.2 Configuration Files', level=2)
    doc.add_paragraph('config.py - Contains all configuration settings including MongoDB connection, JWT settings, and file upload configurations.')
    
    config_code = """import os
from dotenv import load_dotenv

load_dotenv()

# MongoDB Configuration
MONGODB_URL = os.getenv("MONGODB_URL", "mongodb://localhost:27017")
DATABASE_NAME = os.getenv("DATABASE_NAME", "instahub")

# JWT Configuration
SECRET_KEY = os.getenv("SECRET_KEY", "your-secret-key-change-this")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_HOURS = 24

# File Upload Configuration
UPLOAD_DIR = "uploads"
IMAGES_DIR = f"{UPLOAD_DIR}/images"
PROFILES_DIR = f"{UPLOAD_DIR}/profiles"
ALLOWED_IMAGE_EXTENSIONS = ['.jpg', '.jpeg', '.png', '.gif']
MAX_IMAGE_SIZE = (1080, 1080)"""
    
    config_para = doc.add_paragraph(config_code)
    config_para.style = 'No Spacing'
    for run in config_para.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    
    doc.add_paragraph('database.py - Handles MongoDB connection and provides database collections.')
    
    db_code = """from pymongo import MongoClient
from config import MONGODB_URL, DATABASE_NAME

client = MongoClient(MONGODB_URL)
db = client[DATABASE_NAME]

# Collections
users_collection = db.users
posts_collection = db.posts
comments_collection = db.comments
follows_collection = db.follows
notifications_collection = db.notifications"""
    
    db_para = doc.add_paragraph(db_code)
    db_para.style = 'No Spacing'
    for run in db_para.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
    
    # Auth section
    auth_heading = doc.add_heading('4.3 Authentication Module', level=2)
    doc.add_paragraph('auth.py - Handles JWT token generation, password hashing, and user authentication.')
    doc.add_paragraph('This module includes functions for:')
    auth_features = [
        'Password hashing using bcrypt',
        'JWT token creation and validation',
        'Current user dependency for protected routes'
    ]
    for item in auth_features:
        doc.add_paragraph(item, style='List Bullet')
    
    # Controllers section
    controllers_heading = doc.add_heading('4.4 Controllers', level=2)
    doc.add_paragraph('The project follows a controller-based architecture for better code organization:')
    controllers = [
        'auth_controller.py - User registration and login logic',
        'user_controller.py - User profile, follow/unfollow, and search functionality',
        'post_controller.py - Post creation, likes, and comments',
        'feed_controller.py - Feed generation from followed users',
        'notification_controller.py - Notification management'
    ]
    for item in controllers:
        doc.add_paragraph(item, style='List Bullet')
    
    # Frontend section
    frontend_heading = doc.add_heading('4.5 Frontend', level=2)
    doc.add_paragraph('The frontend consists of HTML, CSS, and JavaScript files located in the frontend/ directory:')
    frontend_files = [
        'index.html - Main HTML structure',
        'style.css - Styling and responsive design',
        'app.js - JavaScript for API interactions and UI logic'
    ]
    for item in frontend_files:
        doc.add_paragraph(item, style='List Bullet')
    doc.add_paragraph('[UI Snapshots would be inserted here showing the login page, home feed, profile page, and post creation interface]')
    
    # Page break
    doc.add_page_break()
    
    # Testing
    test_heading = doc.add_heading('TESTING', level=1)
    test_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    test_table = doc.add_table(rows=1, cols=5)
    test_table.style = 'Light Grid Accent 1'
    test_hdr = test_table.rows[0].cells
    test_hdr[0].text = 'Test Case'
    test_hdr[1].text = 'Input'
    test_hdr[2].text = 'Expected Output'
    test_hdr[3].text = 'Actual Result'
    test_hdr[4].text = 'Status'
    for cell in test_hdr:
        cell.paragraphs[0].runs[0].bold = True
    
    test_cases = [
        ('User Registration', 'Valid user credentials', 'User account created successfully', 'Works correctly', 'Passed'),
        ('User Login', 'Valid username and password', 'JWT token generated and returned', 'Working properly', 'Passed'),
        ('Create Post', 'Image file and caption', 'Post created and saved to database', 'Working properly', 'Passed'),
        ('Like/Unlike Post', 'Click like button on post', 'Post like status toggled', 'Working properly', 'Passed'),
        ('Add Comment', 'Comment text on a post', 'Comment added to post', 'Working properly', 'Passed'),
        ('Follow/Unfollow User', 'Click follow button on user profile', 'Follow status toggled', 'Working properly', 'Passed'),
        ('View Feed', 'Navigate to home feed', 'Posts from followed users displayed', 'Working properly', 'Passed'),
        ('Image Upload Validation', 'Invalid file type or large file', 'Error message displayed', 'Validation working', 'Passed'),
        ('Responsive Design', 'Open on different devices', 'Layout adjusts appropriately', 'Responsive design works', 'Passed'),
    ]
    
    for case, inp, exp, act, stat in test_cases:
        row = test_table.add_row().cells
        row[0].text = case
        row[1].text = inp
        row[2].text = exp
        row[3].text = act
        row[4].text = stat
    
    # Page break
    doc.add_page_break()
    
    # Future Enhancement
    future_heading = doc.add_heading('FUTURE ENHANCEMENT', level=1)
    future_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    enhancements = [
        'Add real-time messaging and chat functionality between users',
        'Implement story feature similar to Instagram Stories',
        'Add video upload and playback support',
        'Implement direct messaging (DM) system',
        'Add post editing and deletion capabilities',
        'Implement advanced search with filters and hashtags',
        'Add explore/discover page with trending posts',
        'Implement push notifications for better user engagement',
        'Add multiple image upload support for single post',
        'Implement image filters and editing tools',
        'Add user profile customization options',
        'Implement activity log and analytics for users',
        'Add post sharing functionality',
        'Implement user blocking and reporting features',
        'Add support for multiple languages',
        'Implement cloud storage integration (AWS S3, Google Cloud Storage)',
        'Add email verification and password reset functionality',
        'Implement two-factor authentication (2FA) for enhanced security'
    ]
    
    for item in enhancements:
        doc.add_paragraph(item, style='List Bullet')
    
    # Page break
    doc.add_page_break()
    
    # Conclusion
    concl_heading = doc.add_heading('CONCLUSION', level=1)
    concl_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph(
        'The InstaHub project has been an invaluable learning experience that helped me understand how modern web applications are built from the ground up. The main goal of this project was to create a functional social media platform MVP that demonstrates core features like user authentication, content sharing, and social interactions, and it successfully achieves that objective.'
    )
    
    doc.add_paragraph(
        'Instead of struggling with complex frameworks and architectures, users can interact with a clean, intuitive interface to share moments, connect with others, and engage through likes, comments, and follows. The project demonstrates a complete full-stack application using FastAPI for the backend, MongoDB for data storage, and modern web technologies for the frontend.'
    )
    
    doc.add_paragraph(
        'While working on this project, I learned how different technologies come together to build a complete application. I gained hands-on experience with RESTful API design, database management with MongoDB, JWT-based authentication, file handling, and frontend development. I also understood the importance of proper code organization, error handling, and testing.'
    )
    
    doc.add_paragraph(
        'The system worked well during testing and was able to handle user registration, post creation, social interactions, and feed generation correctly based on the information provided by users. The modular controller-based architecture makes the codebase maintainable and scalable for future enhancements.'
    )
    
    doc.add_paragraph(
        'Overall, this project was a meaningful learning experience. It not only improved my technical skills in full-stack development but also helped me think from a user\'s perspective, considering usability and user experience. In the future, InstaHub can be improved further by adding features like real-time messaging, video support, advanced search, and cloud storage integration, making it even more useful and engaging for users.'
    )
    
    # Save document
    doc.save('MinorProject.docx')
    print("Document created successfully: MinorProject.docx")

if __name__ == '__main__':
    create_documentation()

